"""Render a YAML deck spec into a .pptx in the course's house style.

The style follows the LangGraph deck (05-graph-engineering/sources/langgraph_extended.pptx):
navy title bar with an orange edge, a pale breadcrumb strip, white cards with coloured
stripes, square bullet markers, and a tinted takeaway strip.

    uv run -q --with python-pptx --with pyyaml --with pillow --with resvg-py \
        python scripts/slides/build_deck.py ai-engineering/<pillar>/slides/deck.yaml [--preview DIR] [--pdf] [--strict]

The spec format and layout catalogue are in .claude/skills/slides/references/deck-spec.md.
Text is shrunk to fit its region; anything still too big at the minimum size is reported
as a warning (and fails the run under --strict).
"""

from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import yaml
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# ---------------------------------------------------------------------------
# Theme
# ---------------------------------------------------------------------------

FONT = "Helvetica Neue"
MONO = "Menlo"

NAVY = "1E3A5F"
ORANGE = "D4621A"
TEAL = "0EA5C9"
TEAL_DARK = "0F7490"
PURPLE = "5B21B6"
GREEN = "166534"
BG = "F4F6F9"
PANEL = "E8EDF5"
TINT = "E8F4F8"
INK = "1E293B"
SLATE = "475569"
MUTED = "94A3B8"
LIGHT = "CBD5E1"
WHITE = "FFFFFF"
ROW_ALT = "F8FAFC"
CODE_BG = "0F172A"
CODE_FG = "E2E8F0"
BORDER = "D5DEEA"

PART_ACCENTS = [ORANGE, NAVY, TEAL_DARK, PURPLE, TEAL, GREEN]
FLOW_FILLS = [NAVY, TEAL_DARK, TEAL, PURPLE, GREEN]

SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.55
CONTENT_W = SLIDE_W - 2 * MARGIN
TITLE_BAR_H = 0.95
CRUMB_H = 0.36
CONTENT_TOP = TITLE_BAR_H + CRUMB_H + 0.3
CONTENT_BOTTOM = 6.8
TAKEAWAY_H = 0.55
FOOTER_Y = 7.02

WARNINGS: list[str] = []


def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


def is_light(hex_: str) -> bool:
    r, g, b = (int(hex_[i:i + 2], 16) for i in (0, 2, 4))
    return 0.299 * r + 0.587 * g + 0.114 * b > 170


# ---------------------------------------------------------------------------
# Text measurement (estimates; Helvetica Neue averages ~0.52 em per char)
# ---------------------------------------------------------------------------

MARKUP = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def plain(s: str) -> str:
    return str(s).replace("**", "").replace("`", "")


def est_lines(s: str, width_in: float, pt: float, mono: bool = False, bold: bool = False) -> int:
    em = 0.61 if mono else (0.56 if bold else 0.52)
    per_line = max(1, int(width_in / (pt / 72 * em)))
    total = 0
    for para in plain(s).split("\n"):
        if mono:
            total += max(1, -(-len(para) // per_line))
            continue
        lines, cur = 1, 0
        for word in para.split(" "):
            n = len(word)
            if cur == 0:
                cur = n
            elif cur + 1 + n <= per_line:
                cur += 1 + n
            else:
                lines, cur = lines + 1, n
        total += lines
    return total


def est_h(s: str, width_in: float, pt: float, spacing: float = 1.1, **kw) -> float:
    return est_lines(s, width_in, pt, **kw) * pt * 1.2 * spacing / 72


def warn(ctx: str, msg: str) -> None:
    WARNINGS.append(f"{ctx}: {msg}")


# ---------------------------------------------------------------------------
# Drawing primitives
# ---------------------------------------------------------------------------


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.shadow.inherit = False
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(0.75)
    else:
        shp.line.fill.background()
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius
    return shp


def add_runs(p, s: str, size, color, bold=False, italic=False, font=FONT, code_color=TEAL_DARK):
    for chunk in MARKUP.split(str(s)):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            # bold may contain `code`; recurse with bold on and the emphasis colour
            add_runs(p, chunk[2:-2], size, color if color not in (SLATE, MUTED) else INK, bold=True,
                     italic=italic, font=font, code_color=code_color)
            continue
        r = p.add_run()
        f = r.font
        f.size = Pt(size)
        f.italic = italic
        if chunk.startswith("`") and chunk.endswith("`"):
            r.text = chunk[1:-1]
            f.bold, f.name = bold, MONO
            f.size = Pt(size * 0.9)
            f.color.rgb = rgb(code_color)
        else:
            r.text = chunk
            f.bold, f.name = bold, font
            f.color.rgb = rgb(color)


def text(slide, x, y, w, h, content, size, color=INK, bold=False, italic=False, font=FONT,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.1, code_color=TEAL_DARK):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = content if isinstance(content, list) else str(content).split("\n")
    for i, ptext in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        add_runs(p, ptext, size, color, bold=bold, italic=italic, font=font, code_color=code_color)
    return tb


def line(slide, x1, y1, x2, y2, color, width=1.5, arrow=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    if arrow:
        ln = c.line._get_or_add_ln()
        tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
        ln.append(tail)
    return c


def set_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


# ---------------------------------------------------------------------------
# Composite blocks
# ---------------------------------------------------------------------------


def norm_items(items) -> list[dict]:
    out = []
    for it in items or []:
        if isinstance(it, dict):
            out.append({"text": it.get("text", ""), "sub": it.get("sub", []) or []})
        else:
            out.append({"text": str(it), "sub": []})
    return out


def bullet_layout(items, w, h, start, minimum):
    """Pick the largest font size at which the bullet list fits in (w, h)."""
    items = norm_items(items)
    for size in range(start, minimum - 1, -1):
        sub_size = max(size - 3, 11)
        gap = size * 0.6 / 72
        total = 0.0
        for it in items:
            total += est_h(it["text"], w - 0.35, size)
            for s in it["sub"]:
                total += est_h(s, w - 0.7, sub_size) + gap * 0.4
            total += gap
        total -= gap if items else 0
        if total <= h:
            return items, size, True
    return items, minimum, False


def bullets(slide, x, y, w, h, items, ctx, marker=ORANGE, color=INK, start=20, minimum=13,
            anchor_top=True):
    items, size, ok = bullet_layout(items, w, h, start, minimum)
    if not ok:
        warn(ctx, f"bullets overflow {w:.1f}x{h:.1f}in even at {minimum}pt — cut text")
    sub_size = max(size - 3, 11)
    gap = size * 0.6 / 72
    sq = size / 72 * 0.42
    cy = y
    for it in items:
        th = est_h(it["text"], w - 0.35, size)
        rect(slide, x, cy + size / 72 * 0.38, sq, sq, marker)
        text(slide, x + 0.35, cy, w - 0.35, th + 0.1, it["text"], size, color)
        cy += th
        for s in it["sub"]:
            cy += gap * 0.4
            sh = est_h(s, w - 0.7, sub_size)
            text(slide, x + 0.55, cy, 0.2, sh, "–", sub_size, MUTED)
            text(slide, x + 0.75, cy, w - 0.75, sh + 0.1, s, sub_size, SLATE)
            cy += sh
        cy += gap
    return cy - gap


def card(slide, x, y, w, h, spec, accent, ctx, header_style="stripe", body_start=17, body_min=12):
    """White card with an accent stripe (or a solid header band) + title + body/bullets + tag."""
    rect(slide, x, y, w, h, WHITE, line=BORDER)
    title = spec.get("title", "")
    pad = 0.25
    inner_w = w - 2 * pad
    title_size = 18 if inner_w > 4 else 16
    if header_style == "band":
        th = est_h(title, inner_w, title_size, bold=True) + 0.28
        rect(slide, x, y, w, th, accent)
        on_dark = accent not in (PANEL, TINT, BG)
        text(slide, x + pad, y, inner_w, th, title, title_size, WHITE if on_dark else SLATE, bold=True,
             anchor=MSO_ANCHOR.MIDDLE)
        cy = y + th + 0.22
    else:
        rect(slide, x, y, 0.09, h, accent)
        th = est_h(title, inner_w, title_size, bold=True) if title else 0
        if title:
            text(slide, x + pad + 0.05, y + 0.2, inner_w, th + 0.1, title, title_size, INK, bold=True)
        cy = y + 0.2 + th + (0.18 if title else 0)
    tag = spec.get("tag")
    tag_h = 0.45 if tag else 0
    body_bottom = y + h - 0.2 - tag_h
    bx = x + pad + (0.05 if header_style == "stripe" else 0)
    if spec.get("bullets"):
        bullets(slide, bx, cy, inner_w - 0.05, body_bottom - cy, spec["bullets"], ctx,
                marker=accent if accent not in (PANEL, TINT) else MUTED, start=body_start, minimum=body_min)
    elif spec.get("body"):
        body = spec["body"]
        for size in range(body_start, body_min - 1, -1):
            if est_h(body, inner_w, size) <= body_bottom - cy:
                break
        else:
            warn(ctx, f"card '{plain(title)}' body overflows")
        text(slide, bx, cy, inner_w - 0.05, body_bottom - cy, body, size, SLATE)
    elif spec.get("code"):
        code_block(slide, bx, cy, inner_w - 0.05, body_bottom - cy, spec["code"], ctx, start=14, minimum=10)
    if tag:
        rect(slide, x + 0.2, y + h - 0.2 - 0.38, w - 0.4, 0.38, PANEL)
        text(slide, x + 0.3, y + h - 0.2 - 0.38, w - 0.6, 0.38, tag, 12, SLATE, anchor=MSO_ANCHOR.MIDDLE,
             align=PP_ALIGN.CENTER)


def code_block(slide, x, y, w, h, code, ctx, lang=None, start=16, minimum=10):
    code = str(code).rstrip("\n")
    pad = 0.25
    top_pad = 0.45 if lang else pad
    for size in range(start, minimum - 1, -1):
        need = est_h(code, w - 2 * pad, size, spacing=1.0, mono=True)
        if need <= h - top_pad - pad:
            break
    else:
        warn(ctx, f"code block overflows {w:.1f}x{h:.1f}in even at {minimum}pt")
    need = est_h(code, w - 2 * pad, size, spacing=1.0, mono=True)
    box_h = min(h, need + top_pad + pad + 0.05)
    rect(slide, x, y, w, box_h, CODE_BG, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.03)
    if lang:
        text(slide, x + pad, y + 0.12, w - 2 * pad, 0.3, lang.upper(), 10, MUTED, bold=True)
    text(slide, x + pad, y + top_pad, w - 2 * pad, box_h - top_pad - pad, code.split("\n"), size, CODE_FG,
         font=MONO, spacing=1.0)
    return y + box_h


# ---------------------------------------------------------------------------
# Slide chrome
# ---------------------------------------------------------------------------


def chrome(slide, spec, ctx_info):
    set_bg(slide, BG)
    rect(slide, 0, 0, SLIDE_W, TITLE_BAR_H, NAVY)
    rect(slide, 0, 0, 0.14, TITLE_BAR_H, ctx_info["accent"])
    title = spec.get("title", "")
    kicker = spec.get("kicker")
    title_w = CONTENT_W - (3.6 if kicker else 0)
    size = 28
    while size > 20 and est_lines(title, title_w, size, bold=True) > 1:
        size -= 1
    text(slide, MARGIN, 0, title_w, TITLE_BAR_H, title, size, WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    if kicker:
        ksize = 14
        while ksize > 11 and est_lines(kicker, 3.5, ksize) > 1:
            ksize -= 1
        if est_lines(kicker, 3.5, ksize) > 1:
            warn(ctx_info["ctx"], f"kicker '{kicker}' wraps — keep it under ~40 characters")
        text(slide, SLIDE_W - MARGIN - 3.5, 0, 3.5, TITLE_BAR_H, kicker, ksize, ctx_info["kicker_color"],
             align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    rect(slide, 0, TITLE_BAR_H, SLIDE_W, CRUMB_H, PANEL)
    text(slide, MARGIN, TITLE_BAR_H, 8, CRUMB_H, ctx_info["crumb"], 11, MUTED, anchor=MSO_ANCHOR.MIDDLE)
    # footer
    if ctx_info.get("source"):
        text(slide, MARGIN, FOOTER_Y, 9, 0.3, ctx_info["source"], 10, MUTED, italic=True)
    text(slide, SLIDE_W - MARGIN - 1.5, FOOTER_Y, 1.5, 0.3, str(ctx_info["number"]), 10, MUTED,
         align=PP_ALIGN.RIGHT)

    top = CONTENT_TOP
    lead = spec.get("lead")
    if lead:
        lh = est_h(lead, CONTENT_W, 17)
        text(slide, MARGIN, top - 0.05, CONTENT_W, lh + 0.1, lead, 17, SLATE)
        top += lh + 0.25
    bottom = CONTENT_BOTTOM
    tk = spec.get("takeaway")
    if tk:
        size = 16
        while size > 12 and est_lines(tk, CONTENT_W - 0.5, size) > 1:
            size -= 1
        th = max(TAKEAWAY_H, est_h(tk, CONTENT_W - 0.5, size) + 0.25)
        ty = CONTENT_BOTTOM - th
        rect(slide, MARGIN, ty, CONTENT_W, th, TINT)
        rect(slide, MARGIN, ty, 0.08, th, ctx_info["accent"])
        text(slide, MARGIN + 0.3, ty, CONTENT_W - 0.5, th, tk, size, NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE,
             code_color=NAVY)
        bottom = ty - 0.25
    if bottom - top < 1.5:
        warn(ctx_info["ctx"], f"only {bottom - top:.1f}in left for content — shorten the lead/takeaway")
    return top, bottom


# ---------------------------------------------------------------------------
# Layouts
# ---------------------------------------------------------------------------


def lay_bullets(slide, spec, top, bottom, ctx, info):
    aside = spec.get("aside")
    w = CONTENT_W if not aside else CONTENT_W * 0.56
    bullets(slide, MARGIN, top, w, bottom - top, spec.get("bullets", []), ctx, marker=info["accent"])
    if aside:
        ax = MARGIN + w + 0.4
        aw = CONTENT_W - w - 0.4
        if aside.get("code") and not aside.get("title"):
            code_block(slide, ax, top, aw, bottom - top, aside["code"], ctx, lang=aside.get("lang"))
        else:
            card(slide, ax, top, aw, bottom - top, aside, aside.get("accent", TEAL_DARK), ctx,
                 header_style="band")


def lay_compare(slide, spec, top, bottom, ctx, info):
    left, right = spec["left"], spec["right"]
    emph = spec.get("emphasis", "right")
    gap = 0.4
    w = (CONTENT_W - gap) / 2
    lacc = PANEL if emph == "right" else NAVY
    racc = NAVY if emph in ("right", "none") else PANEL
    if emph == "none":
        lacc = TEAL_DARK
    card(slide, MARGIN, top, w, bottom - top, left, left.get("accent", lacc), ctx, header_style="band")
    card(slide, MARGIN + w + gap, top, w, bottom - top, right, right.get("accent", racc), ctx, header_style="band")


def lay_cards(slide, spec, top, bottom, ctx, info):
    cards = spec["cards"]
    n = len(cards)
    cols = n if n <= 4 else 3
    rows = -(-n // cols)
    gap = 0.3
    w = (CONTENT_W - gap * (cols - 1)) / cols
    h = (bottom - top - gap * (rows - 1)) / rows
    start = 20 if cols <= 3 else 18
    # Body-only cards: take the largest font that fits, then shrink the cards to their content so a
    # short body doesn't sit in a tall empty card.
    if all(c.get("body") and not c.get("bullets") and not c.get("tag") for c in cards):
        inner_w = w - 0.5
        title_size = 18 if inner_w > 4 else 16

        def need(size):
            return max(est_h(c.get("title", ""), inner_w, title_size, bold=True) + 0.38
                       + est_h(c["body"], inner_w, size) + 0.35 for c in cards)

        while start > 13 and need(start) > h:
            start -= 1
        h = min(h, need(start) + 0.15)
    for i, c in enumerate(cards):
        r, col = divmod(i, cols)
        acc = c.get("accent") or PART_ACCENTS[i % len(PART_ACCENTS)]
        card(slide, MARGIN + col * (w + gap), top + r * (h + gap), w, h, c, acc, ctx, body_start=start)


def lay_table(slide, spec, top, bottom, ctx, info):
    header = spec["header"]
    rows = spec["rows"]
    ncol = len(header)
    for r in rows:
        if len(r) != ncol:
            sys.exit(f"ERROR slide '{spec.get('title')}': row has {len(r)} cells, header has {ncol} "
                     f"(quote cells that contain commas): {r}")
    ratios = spec.get("widths") or [1] * ncol
    tot = sum(ratios)
    widths = [CONTENT_W * r / tot for r in ratios]
    highlight = set(spec.get("highlight", []))
    avail = bottom - top
    cell_pad = 0.12

    def row_h(cells, size, bold=False):
        return max(est_h(c, wd - 2 * cell_pad, size, spacing=1.0, bold=bold) for c, wd in zip(cells, widths)) + 0.16

    for size in range(17, 10, -1):
        heights = [row_h(header, size, True)] + [row_h(r, size) for r in rows]
        if sum(heights) <= avail:
            break
    else:
        warn(ctx, "table overflows even at 11pt — split it or trim cells")
    shape = slide.shapes.add_table(len(rows) + 1, ncol, Inches(MARGIN), Inches(top), Inches(CONTENT_W),
                                   Inches(sum(heights)))
    tbl = shape.table
    tblPr = tbl._tbl.tblPr
    style = tblPr.find(qn("a:tableStyleId"))
    if style is None:
        style = tblPr.makeelement(qn("a:tableStyleId"), {})
        tblPr.append(style)
    style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # "No Style, Table Grid"
    for j, wd in enumerate(widths):
        tbl.columns[j].width = Inches(wd)
    for i, cells in enumerate([header] + rows):
        tbl.rows[i].height = Inches(heights[i])
        for j, val in enumerate(cells):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(cell_pad)
            cell.margin_top = cell.margin_bottom = Inches(0.06)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if i == 0:
                cell.fill.fore_color.rgb = rgb(NAVY)
            elif (i - 1) in highlight:
                cell.fill.fore_color.rgb = rgb(TINT)
            else:
                cell.fill.fore_color.rgb = rgb(WHITE if i % 2 else ROW_ALT)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.line_spacing = 1.0
            if i == 0:
                add_runs(p, val, size, WHITE, bold=True, code_color=TEAL)
            else:
                add_runs(p, val, size, INK if j == 0 else SLATE, bold=(j == 0 and spec.get("bold_first", True)))


def lay_flow(slide, spec, top, bottom, ctx, info):
    steps = spec["steps"]
    n = len(steps)
    gap = 0.55
    box_w = min(2.6, (CONTENT_W - gap * (n - 1)) / n)
    total_w = box_w * n + gap * (n - 1)
    x0 = MARGIN + (CONTENT_W - total_w) / 2
    lab_size = 17 if box_w > 2 else 14
    sub_size = lab_size - 4
    box_h = max(
        est_h(plain(s["label"] if isinstance(s, dict) else s), box_w - 0.3, lab_size, bold=True)
        + (est_h(s.get("sub", ""), box_w - 0.3, sub_size) + 0.08 if isinstance(s, dict) and s.get("sub") else 0)
        for s in steps
    ) + 0.45
    y = top + 0.1
    for i, s in enumerate(steps):
        s = s if isinstance(s, dict) else {"label": s}
        x = x0 + i * (box_w + gap)
        fill = s.get("accent") or (ORANGE if i == n - 1 and spec.get("end_accent", True) else FLOW_FILLS[i % len(FLOW_FILLS)])
        rect(slide, x, y, box_w, box_h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
        content = [s["label"]] + ([s["sub"]] if s.get("sub") else [])
        light = is_light(fill)
        tb = text(slide, x + 0.15, y, box_w - 0.3, box_h, content, lab_size, INK if light else WHITE,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, code_color=TEAL_DARK if light else TINT)
        for r in tb.text_frame.paragraphs[0].runs:
            r.font.bold = True
        if s.get("sub"):
            for r in tb.text_frame.paragraphs[1].runs:
                r.font.size = Pt(sub_size)
                r.font.color.rgb = rgb(SLATE if light else TINT)
        if i < n - 1:
            line(slide, x + box_w + 0.06, y + box_h / 2, x + box_w + gap - 0.06, y + box_h / 2, SLATE, 2, arrow=True)
    cy = y + box_h
    loop = spec.get("loop")
    if loop:
        lx1 = x0 + total_w - box_w / 2
        lx0 = x0 + box_w / 2
        ly = cy + 0.45
        line(slide, lx1, cy, lx1, ly, GREEN, 1.75)
        line(slide, lx1, ly, lx0, ly, GREEN, 1.75)
        line(slide, lx0, ly, lx0, cy + 0.04, GREEN, 1.75, arrow=True)
        text(slide, lx0, ly + 0.04, lx1 - lx0, 0.35, loop, 13, GREEN, italic=True, align=PP_ALIGN.CENTER)
        cy = ly + 0.45
    if spec.get("bullets"):
        bullets(slide, MARGIN, cy + 0.35, CONTENT_W, bottom - cy - 0.35, spec["bullets"], ctx,
                marker=info["accent"], start=19)
    elif spec.get("caption"):
        text(slide, MARGIN, cy + 0.35, CONTENT_W, bottom - cy - 0.35, spec["caption"], 17, SLATE,
             align=PP_ALIGN.CENTER)


def lay_steps(slide, spec, top, bottom, ctx, info):
    items = spec["steps"]
    n = len(items)
    axis = spec.get("axis")
    axis_w = 1.25 if axis else 0
    x = MARGIN + axis_w
    w = CONTENT_W - axis_w
    gap = 0.12
    row_h = (bottom - top - gap * (n - 1)) / n
    num_d = min(0.5, row_h - 0.12)
    for size in range(18, 11, -1):
        fits = all(
            est_h(f"**{it['title']}** — {it.get('text', '')}" if it.get("text") else it["title"], w - num_d - 0.6,
                  size) <= row_h - 0.08
            for it in items
        )
        if fits:
            break
    else:
        warn(ctx, "steps overflow — shorten the step text")
    for i, it in enumerate(items):
        y = top + i * (row_h + gap)
        acc = it.get("accent") or PART_ACCENTS[i % len(PART_ACCENTS)]
        rect(slide, x, y, w, row_h, WHITE, line=BORDER)
        rect(slide, x + 0.18, y + (row_h - num_d) / 2, num_d, num_d, acc, shape=MSO_SHAPE.OVAL)
        text(slide, x + 0.18, y + (row_h - num_d) / 2, num_d, num_d, str(it.get("n", i + 1)), size - 2, WHITE,
             bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        body = f"**{it['title']}** — {it['text']}" if it.get("text") else f"**{it['title']}**"
        text(slide, x + num_d + 0.4, y, w - num_d - 0.6, row_h, body, size, SLATE, anchor=MSO_ANCHOR.MIDDLE)
    if axis:
        # "weaker → stronger guarantee": first half labels the top of the arrow, second half the bottom.
        # (Rotated text is avoided: Quick Look and some viewers ignore the rotation.)
        head, _, tail = (p.strip() for p in str(axis).partition("→"))
        lab_w = axis_w - 0.15
        cx = MARGIN + lab_w / 2
        lab_h = 0.55
        text(slide, MARGIN, top, lab_w, lab_h, head, 12, MUTED, italic=True, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.BOTTOM)
        if tail:
            text(slide, MARGIN, bottom - lab_h, lab_w, lab_h, tail, 12, MUTED, italic=True, align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.TOP)
        line(slide, cx, top + lab_h + 0.1, cx, bottom - (lab_h + 0.1 if tail else 0.1), MUTED, 2.25, arrow=True)


def lay_code(slide, spec, top, bottom, ctx, info):
    side = spec.get("bullets")
    w = CONTENT_W if not side else CONTENT_W * 0.6
    code_block(slide, MARGIN, top, w, bottom - top, spec["code"], ctx, lang=spec.get("lang"))
    if side:
        bullets(slide, MARGIN + w + 0.4, top + 0.05, CONTENT_W - w - 0.4, bottom - top, side, ctx,
                marker=info["accent"], start=18)


def lay_statement(slide, spec, top, bottom, ctx, info):
    stmt = spec["statement"]
    size = 36
    w = CONTENT_W - 1.2
    while size > 22 and est_h(stmt, w, size, bold=True) > (bottom - top) * 0.7:
        size -= 2
    h = est_h(stmt, w, size, bold=True)
    by = top + (bottom - top - h) / 2 - 0.3
    rect(slide, MARGIN + 0.3, by, 0.12, h + (0.6 if spec.get("attribution") else 0.1), info["accent"])
    text(slide, MARGIN + 0.8, by, w, h + 0.2, stmt, size, NAVY, bold=True, spacing=1.05, code_color=TEAL_DARK)
    if spec.get("attribution"):
        text(slide, MARGIN + 0.8, by + h + 0.15, w, 0.4, spec["attribution"], 17, SLATE, italic=True)


def raster(path: Path) -> Path:
    """python-pptx only embeds bitmaps, so render an SVG to a PNG beside the temp files."""
    if path.suffix.lower() != ".svg":
        return path
    import resvg_py

    out = Path(tempfile.gettempdir()) / "build_deck_svg" / f"{path.stem}-{abs(hash(str(path)))}.png"
    if not out.exists() or out.stat().st_mtime < path.stat().st_mtime:
        out.parent.mkdir(exist_ok=True)
        out.write_bytes(bytes(resvg_py.svg_to_bytes(svg_path=str(path), width=2000, background="white")))
    return out


def shrink(path: Path, w_in: float, h_in: float, dpi: int = 220) -> Path:
    """Downsample to the size the image is shown at, so a deck stays small enough to commit."""
    from PIL import Image

    img = Image.open(path)
    target = (round(w_in * dpi), round(h_in * dpi))
    if img.width <= target[0] and path.stat().st_size < 150_000:
        return path
    out_dir = Path(tempfile.gettempdir()) / "build_deck_img"
    out_dir.mkdir(exist_ok=True)
    stem = f"{path.stem}-{abs(hash((str(path), target)))}"
    if img.width > target[0]:
        img = img.convert("RGBA" if img.mode in ("RGBA", "LA", "P") else "RGB").resize(target, Image.LANCZOS)
    png = out_dir / f"{stem}.png"
    img.save(png, optimize=True)
    # Screenshots and photos compress far better as JPEG; keep PNG when it is already smaller.
    jpg = out_dir / f"{stem}.jpg"
    flat = Image.new("RGB", img.size, "white")
    flat.paste(img, mask=img.getchannel("A") if img.mode == "RGBA" else None)
    flat.save(jpg, quality=88, optimize=True)
    best = min((png, jpg, path), key=lambda p: p.stat().st_size)
    return best


def lay_image(slide, spec, top, bottom, ctx, info):
    from PIL import Image

    path = (info["base"] / spec["image"]).resolve()
    side = spec.get("bullets")
    region_w = CONTENT_W if not side else CONTENT_W * 0.55
    cap = spec.get("caption")
    region_h = bottom - top - (0.45 if cap else 0)
    if not path.exists():
        warn(ctx, f"image not found: {path}")
        return
    path = raster(path)
    iw, ih = Image.open(path).size
    scale = min(region_w / iw, region_h / ih)
    w, h = iw * scale, ih * scale
    x = MARGIN + (region_w - w) / 2
    rect(slide, x - 0.06, top - 0.06, w + 0.12, h + 0.12, WHITE, line=BORDER)
    slide.shapes.add_picture(str(shrink(path, w, h)), Inches(x), Inches(top), Inches(w), Inches(h))
    if cap:
        text(slide, MARGIN, top + h + 0.15, region_w, 0.4, cap, 12, MUTED, italic=True, align=PP_ALIGN.CENTER)
    if side:
        bullets(slide, MARGIN + region_w + 0.4, top, CONTENT_W - region_w - 0.4, bottom - top, side, ctx,
                marker=info["accent"], start=18)


LAYOUTS = {
    "bullets": lay_bullets,
    "compare": lay_compare,
    "cards": lay_cards,
    "table": lay_table,
    "flow": lay_flow,
    "steps": lay_steps,
    "code": lay_code,
    "statement": lay_statement,
    "image": lay_image,
}


# ---------------------------------------------------------------------------
# Title and section slides
# ---------------------------------------------------------------------------


def title_slide(prs, deck, parts):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s, BG)
    rect(s, 8.2, -1.2, 5.6, 5.6, "DCEFF7", shape=MSO_SHAPE.OVAL)
    rect(s, 9.6, 4.4, 4.4, 4.4, "F6E6DB", shape=MSO_SHAPE.OVAL)
    panel_w = 6.4
    rect(s, 0, 0, panel_w, SLIDE_H, NAVY)
    rect(s, 0, 0, 0.18, SLIDE_H, ORANGE)
    title = deck["title"]
    size = 60
    while size > 36 and est_h(title, panel_w - 1.0, size, bold=True) > 2.2:
        size -= 2
    th = est_h(title, panel_w - 1.0, size, bold=True, spacing=0.95)
    text(s, 0.6, 1.1, panel_w - 1.0, th + 0.2, title, size, WHITE, bold=True, spacing=0.95)
    y = 1.1 + th + 0.35
    if deck.get("subtitle"):
        sh = est_h(deck["subtitle"], panel_w - 1.0, 24)
        text(s, 0.6, y, panel_w - 1.0, sh + 0.1, deck["subtitle"], 24, TEAL)
        y += sh + 0.35
    rect(s, 0.6, y, panel_w - 1.4, 0.05, ORANGE)
    y += 0.35
    if deck.get("tagline"):
        text(s, 0.6, y, panel_w - 1.0, 1.6, deck["tagline"], 15, MUTED, spacing=1.25)
    if deck.get("byline"):
        text(s, 0.6, SLIDE_H - 0.8, panel_w - 1.0, 0.4, deck["byline"], 12, MUTED, italic=True)

    # part cards
    n = len(parts)
    x, w = panel_w + 0.7, SLIDE_W - panel_w - 1.25
    top, bottom = 0.9, SLIDE_H - 0.9
    gap = 0.22
    h = min(1.25, (bottom - top - gap * (n - 1)) / max(n, 1))
    y0 = top + (bottom - top - (h * n + gap * (n - 1))) / 2
    for i, p in enumerate(parts):
        y = y0 + i * (h + gap)
        acc = PART_ACCENTS[i % len(PART_ACCENTS)]
        rect(s, x, y, w, h, WHITE, line=BORDER)
        rect(s, x, y, 0.1, h, acc)
        ts = 17 if h > 0.9 else 14
        text(s, x + 0.3, y + 0.14, w - 0.45, 0.4, f"Part {i + 1}: {p['title']}", ts, INK, bold=True)
        if p.get("summary"):
            text(s, x + 0.3, y + 0.14 + ts / 72 * 1.45, w - 0.45, h - 0.5, p["summary"], ts - 4, SLATE)
    if deck.get("reference"):
        text(s, x, SLIDE_H - 0.55, w, 0.3, deck["reference"], 11, MUTED, italic=True, align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = deck.get("notes", "")


def section_slide(prs, part, idx, number):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    acc = PART_ACCENTS[idx % len(PART_ACCENTS)]
    set_bg(s, NAVY)
    rect(s, 0, 0, 0.18, SLIDE_H, acc if acc != NAVY else TEAL)
    text(s, 0.9, 1.6, 6, 0.5, f"PART {idx + 1}", 18, ORANGE if acc == NAVY else acc if acc not in (NAVY, PURPLE) else TEAL,
         bold=True)
    th = est_h(part["title"], 6.0, 50, bold=True)
    text(s, 0.9, 2.15, 6.0, th + 0.2, part["title"], 50, WHITE, bold=True, spacing=0.95)
    y = 2.15 + th + 0.3
    if part.get("summary"):
        text(s, 0.9, y, 5.8, 1.5, part["summary"], 20, TEAL)
    titles = [plain(sl.get("title", "")) for sl in part.get("slides", []) if sl.get("title")]
    if titles:
        x, w = 7.4, SLIDE_W - 7.4 - 0.8
        avail = SLIDE_H - 2.4

        def list_h(size):
            return sum(est_h(t, w, size, spacing=1.0) for t in titles) + size / 72 * 0.7 * (len(titles) - 1)

        size = 17
        while size > 11 and list_h(size) > avail:
            size -= 1
        total = list_h(size)
        y0 = (SLIDE_H - total) / 2
        line(s, x - 0.35, y0, x - 0.35, y0 + total, "3B5A80", 1.5)
        tb = text(s, x, y0, w, total + 0.2, titles, size, LIGHT, spacing=1.0)
        for p in tb.text_frame.paragraphs[1:]:
            p.space_before = Pt(size * 0.7)
    text(s, SLIDE_W - MARGIN - 1.5, FOOTER_Y, 1.5, 0.3, str(number), 10, "64748B", align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = part.get("notes", "")


# ---------------------------------------------------------------------------
# Build
# ---------------------------------------------------------------------------


def build(spec_path: Path) -> Path:
    spec = yaml.safe_load(spec_path.read_text())
    deck, parts = spec["deck"], spec["parts"]
    base = spec_path.parent

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)

    title_slide(prs, deck, [p for p in parts if p.get("agenda", True)])
    number = 1
    agenda_idx = 0
    for part in parts:
        in_agenda = part.get("agenda", True)
        idx = agenda_idx if in_agenda else len(PART_ACCENTS) - 1
        if in_agenda and part.get("divider", True):
            number += 1
            section_slide(prs, part, idx, number)
        accent = PART_ACCENTS[idx % len(PART_ACCENTS)]
        for sl in part.get("slides", []):
            number += 1
            layout = sl.get("layout", "bullets")
            ctx = f"slide {number} '{plain(sl.get('title', ''))}'"
            if layout not in LAYOUTS:
                warn(ctx, f"unknown layout '{layout}'")
                continue
            s = prs.slides.add_slide(prs.slide_layouts[6])
            crumb = f"Part {agenda_idx + 1} · {part['title']}" if in_agenda else part["title"]
            src = sl.get("source")
            info = {
                "accent": accent if accent != NAVY else ORANGE,
                "kicker_color": ORANGE if accent in (NAVY, ORANGE) else TEAL,
                "crumb": crumb,
                "number": number,
                "source": f"Source: {src}" if src else "",
                "ctx": ctx,
                "base": base,
            }
            top, bottom = chrome(s, sl, info)
            LAYOUTS[layout](s, sl, top, bottom, ctx, info)
            notes = sl.get("notes", "")
            if src:
                notes = f"{notes}\n\nSource: {src}".strip()
            s.notes_slide.notes_text_frame.text = notes
        if in_agenda:
            agenda_idx += 1

    out = base / deck.get("output", f"{base.parent.name}.pptx")
    prs.save(out)
    return out


def render_pngs(pptx: Path, out_dir: Path, size: int = 2400) -> list[Path]:
    """Split the deck into one-slide files and render each to PNG with Quick Look (macOS)."""
    if not shutil.which("qlmanage"):
        print("render: qlmanage not found (macOS only) — skipping", file=sys.stderr)
        return []
    out_dir.mkdir(parents=True, exist_ok=True)
    for old in out_dir.glob("*.png"):
        old.unlink()
    n = len(Presentation(pptx).slides)
    with tempfile.TemporaryDirectory() as tmp:
        singles = []
        for i in range(n):
            prs = Presentation(pptx)
            ids = prs.slides._sldIdLst
            for k, sld in reversed(list(enumerate(ids))):
                if k != i:
                    prs.part.drop_rel(sld.rId)
                    ids.remove(sld)
            # Quick Look hangs indefinitely on a pptx with speaker notes; the render doesn't need them.
            for rId, rel in list(prs.slides[0].part.rels.items()):
                if rel.reltype == RT.NOTES_SLIDE:
                    prs.slides[0].part.drop_rel(rId)
            p = Path(tmp) / f"slide-{i + 1:02d}.pptx"
            prs.save(p)
            singles.append(p)
        try:
            subprocess.run(["qlmanage", "-t", "-s", str(size), "-o", tmp, *map(str, singles)],
                           check=True, capture_output=True, timeout=30 + 5 * n)
        except subprocess.TimeoutExpired:
            print("render: qlmanage timed out — some slides may be missing", file=sys.stderr)
        pngs = []
        for p in singles:
            src = Path(f"{p}.png")
            if src.exists():
                dst = out_dir / f"{p.stem}.png"
                shutil.move(src, dst)
                pngs.append(dst)
        if len(pngs) < n:
            print(f"render: only {len(pngs)} of {n} slides rendered", file=sys.stderr)
    return pngs


def contact_sheets(pngs: list[Path], out_dir: Path) -> list[Path]:
    """Tile the slide PNGs into 2-column x 3-row sheets for quick review."""
    from PIL import Image

    sheets = []
    per = 6
    for sheet_i in range(0, len(pngs), per):
        group = [Image.open(p).convert("RGB") for p in pngs[sheet_i:sheet_i + per]]
        tw = 800
        th = int(group[0].height * tw / group[0].width)
        rows = -(-len(group) // 2)
        sheet = Image.new("RGB", (tw * 2 + 30, th * rows + 15 * (rows + 1)), "white")
        for k, img in enumerate(group):
            r, c = divmod(k, 2)
            sheet.paste(img.resize((tw, th), Image.LANCZOS), (10 + c * (tw + 10), 15 + r * (th + 15)))
        sp = out_dir / f"sheet-{sheet_i // per + 1:02d}.png"
        sheet.save(sp)
        sheets.append(sp)
    return sheets


def write_pdf(pngs: list[Path], pdf: Path, width: int = 1600) -> Path:
    """One page per slide, rasterised from the rendered PNGs (no speaker notes, no selectable text).

    Pages are downsampled to `width` px and JPEG-encoded: sharp on screen, ~5 MB a deck, since
    the PDF is committed and every rebuild adds a copy to git history.
    """
    from PIL import Image

    pages = []
    for p in pngs:
        img = Image.open(p).convert("RGB")
        if img.width > width:
            img = img.resize((width, round(img.height * width / img.width)), Image.LANCZOS)
        pages.append(img)
    # Write beside the target, then rename: a viewer with the PDF open reloads on change and
    # would otherwise read a half-written file and report it corrupt.
    tmp = pdf.with_name(f".{pdf.name}.tmp")
    pages[0].save(tmp, format="PDF", save_all=True, append_images=pages[1:],
                  resolution=width / 13.333, quality=80)
    tmp.replace(pdf)
    return pdf


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("spec", type=Path)
    ap.add_argument("--preview", type=Path, metavar="DIR", help="write per-slide PNGs + contact sheets here")
    ap.add_argument("--pdf", action="store_true", help="also write <output>.pdf next to the .pptx (macOS)")
    ap.add_argument("--strict", action="store_true", help="exit 1 if any slide overflows")
    args = ap.parse_args()

    out = build(args.spec)
    print(f"deck: {out}  ({len(Presentation(out).slides)} slides)")
    if args.preview or args.pdf:
        with tempfile.TemporaryDirectory() as tmp:
            pngs = render_pngs(out, args.preview or Path(tmp))
            if args.preview and pngs:
                for sheet in contact_sheets(pngs, args.preview):
                    print(f"sheet: {sheet}")
            if args.pdf and pngs:
                print(f"pdf: {write_pdf(pngs, out.with_suffix('.pdf'))}")
    for w in WARNINGS:
        print(f"WARN {w}", file=sys.stderr)
    if WARNINGS and args.strict:
        sys.exit(1)


if __name__ == "__main__":
    main()
